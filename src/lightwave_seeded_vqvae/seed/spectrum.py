from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


EXPECTED_POWERS_MW = (0, 5, 10, 15, 20)
EXPECTED_HEADERS = ("λ1 (nm)", "λ2 (nm)", "E1 (dB)", "E2 (dB)")


@dataclass(frozen=True)
class OpticalState:
    pump_nm: int
    power_mw: int
    lambda1_nm: Decimal
    lambda2_nm: Decimal
    extinction1_db: Decimal
    extinction2_db: Decimal
    source_path: str
    source_sha256: str
    source_row: int

    @property
    def state_id(self) -> str:
        return f"{self.pump_nm}nm_{self.power_mw}mW"


def file_sha256(path: str | Path) -> str:
    digest = sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _clean_text(value: object) -> str:
    return unicodedata.normalize("NFC", str(value)).strip()


def _decimal(value: object, *, cell: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{cell} must contain a finite number")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{cell} is not a valid decimal number: {value!r}") from exc
    if not result.is_finite():
        raise ValueError(f"{cell} must contain a finite number")
    return result


def load_optical_workbook(path: str | Path, expected_pump_nm: int | None = None) -> list[OpticalState]:
    workbook_path = Path(path).expanduser().resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError(f"{workbook_path.name}: expected exactly one worksheet")
        sheet = workbook.worksheets[0]
        if sheet.max_row != 6 or sheet.max_column != 5:
            raise ValueError(
                f"{workbook_path.name}: expected A1:E6, got "
                f"{sheet.max_row} rows x {sheet.max_column} columns"
            )

        pump_label = _clean_text(sheet.cell(1, 1).value)
        pump_match = re.fullmatch(r"(\d+)\s*nm", pump_label, flags=re.IGNORECASE)
        if pump_match is None:
            raise ValueError(f"{workbook_path.name}: invalid pump label {pump_label!r}")
        pump_nm = int(pump_match.group(1))
        if expected_pump_nm is not None and pump_nm != expected_pump_nm:
            raise ValueError(
                f"{workbook_path.name}: expected {expected_pump_nm} nm, found {pump_nm} nm"
            )

        headers = tuple(_clean_text(sheet.cell(1, column).value) for column in range(2, 6))
        if headers != EXPECTED_HEADERS:
            raise ValueError(
                f"{workbook_path.name}: expected headers {EXPECTED_HEADERS!r}, found {headers!r}"
            )

        source_hash = file_sha256(workbook_path)
        states: list[OpticalState] = []
        observed_powers: list[int] = []
        for row in range(2, 7):
            power_label = _clean_text(sheet.cell(row, 1).value)
            power_match = re.fullmatch(r"(\d+)\s*mW", power_label, flags=re.IGNORECASE)
            if power_match is None:
                raise ValueError(
                    f"{workbook_path.name}: invalid power label in A{row}: {power_label!r}"
                )
            power_mw = int(power_match.group(1))
            observed_powers.append(power_mw)

            values = [
                _decimal(sheet.cell(row, column).value, cell=f"{sheet.title}!{column}:{row}")
                for column in range(2, 6)
            ]
            lambda1, lambda2, extinction1, extinction2 = values
            if lambda1 >= lambda2:
                raise ValueError(
                    f"{workbook_path.name}: row {row} must satisfy lambda1 < lambda2"
                )
            states.append(
                OpticalState(
                    pump_nm=pump_nm,
                    power_mw=power_mw,
                    lambda1_nm=lambda1,
                    lambda2_nm=lambda2,
                    extinction1_db=extinction1,
                    extinction2_db=extinction2,
                    source_path=str(workbook_path),
                    source_sha256=source_hash,
                    source_row=row,
                )
            )

        if tuple(observed_powers) != EXPECTED_POWERS_MW:
            raise ValueError(
                f"{workbook_path.name}: expected powers {EXPECTED_POWERS_MW}, "
                f"found {tuple(observed_powers)}"
            )
        return states
    finally:
        workbook.close()


def load_optical_states(path_405: str | Path, path_640: str | Path) -> list[OpticalState]:
    states = load_optical_workbook(path_405, expected_pump_nm=405)
    states.extend(load_optical_workbook(path_640, expected_pump_nm=640))
    states.sort(key=lambda state: (state.pump_nm, state.power_mw))
    expected_order = [(pump, power) for pump in (405, 640) for power in EXPECTED_POWERS_MW]
    observed_order = [(state.pump_nm, state.power_mw) for state in states]
    if observed_order != expected_order:
        raise ValueError(f"unexpected state order: {observed_order}")
    return states

